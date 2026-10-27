"""Excel eksport (openpyxl)."""
from __future__ import annotations

from io import BytesIO
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import CustomQuestion, Participant, Subject
from ..utils import fmt_dt, fmt_num, format_phone

_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="D0D7DE")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MEDAL = {1: "FFF2B2", 2: "E8E8E8", 3: "F6D7B0"}


def _sheet(wb: Workbook, title: str, headers: Sequence[str], rows: Iterable[Sequence],
           widths: Sequence[int] | None = None, first: bool = False, places_col: int | None = None):
    ws = wb.active if first else wb.create_sheet()
    # Excel varaq nomi 31 belgidan oshmasligi va ba'zi belgilarsiz bo'lishi kerak
    ws.title = "".join(c for c in title if c not in '[]:*?/\\')[:31] or "Sheet"
    ws.append(list(headers))
    for c in ws[1]:
        c.fill, c.font, c.border = _HEAD_FILL, _HEAD_FONT, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(list(row))
        r = ws.max_row
        for c in ws[r]:
            c.border = _BORDER
            # "=" bilan boshlangan matn formula sifatida bajarilmasin (Excel formula injection)
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
        if places_col is not None:
            place = ws.cell(r, places_col).value
            if place in _MEDAL:
                for c in ws[r]:
                    c.fill = PatternFill("solid", fgColor=_MEDAL[place])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for i, h in enumerate(headers, start=1):
        w = (widths[i - 1] if widths and i - 1 < len(widths) else max(10, min(40, len(str(h)) + 4)))
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def participants_rows(participants: Sequence[Participant], questions: Sequence[CustomQuestion],
                      att: dict | None = None) -> tuple[list, list, list]:
    """Ishtirokchilar jadvali: (sarlavhalar, qatorlar, ustun kengliklari). att — davomat (pid → bool)."""
    headers = ["№", "ID", "Familiya", "Ism", "Telefon", "Maktab", "Sinf", "Fanlar",
               "Manzil", "Instagram", "Telegram", "Til", "Holat", "Ro'yxatdan o'tgan"]
    if att is not None:
        headers.append("Davomat")
    headers += [q.text[:60] for q in questions]
    rows = []
    for i, p in enumerate(participants, start=1):
        ans = {a.question_id: a.value for a in p.answers}
        rows.append([
            i, p.code, p.last_name, p.first_name, format_phone(p.phone),
            p.school.name if p.school else "", p.grade.name if p.grade else "",
            ", ".join(s.name for s in p.subjects), p.address,
            ("@" + p.instagram_username) if p.instagram_username else "",
            ("@" + p.tg_username) if p.tg_username else str(p.telegram_id),
            "Kirill" if p.lang == "cyr" else "Lotin",
            "Faol" if p.status == "active" else "Chetlatilgan",
            fmt_dt(p.created_at),
        ] + ([{True: "Keldi", False: "Kelmadi"}.get(att.get(p.id), "—")] if att is not None else [])
          + [ans.get(q.id, "") for q in questions])
    widths = [6, 8, 18, 16, 18, 26, 8, 30, 34, 16, 16, 8, 12, 18] + ([10] if att is not None else []) \
        + [24] * len(questions)
    return headers, rows, widths


def participants_xlsx(participants: Sequence[Participant], questions: Sequence[CustomQuestion],
                      att: dict | None = None) -> bytes:
    wb = Workbook()
    headers, rows, widths = participants_rows(participants, questions, att)
    _sheet(wb, "Ishtirokchilar", headers, rows, widths=widths, first=True)
    return _to_bytes(wb)


def subject_ranking_table(subject: Subject, rows, show_details: bool = True) -> tuple[list, list]:
    """Fan reytingi jadvali: (sarlavhalar, qatorlar)."""
    import json
    headers = ["O'rin", "ID", "Familiya", "Ism", "Maktab", "Sinf"]
    if show_details and subject.scoring_mode == "groups":
        headers += [f"{g.name} ({g.question_count}×{fmt_num(g.points)})" for g in subject.groups]
    elif show_details and subject.scoring_mode == "uniform":
        headers += [f"To'g'ri ({subject.total_questions} tadan)"]
    headers += ["Ball", "Maks. ball", "Foiz", "Kiritgan hakam"]
    out = []
    for r in rows:
        p = r.participant
        line = [r.place, p.code, p.last_name, p.first_name,
                p.school.name if p.school else "", p.grade.name if p.grade else ""]
        det = json.loads(r.result.details or "{}") if r.result else {}
        if show_details and subject.scoring_mode == "groups":
            line += [det.get(f"g{g.id}", "") for g in subject.groups]
        elif show_details and subject.scoring_mode == "uniform":
            line += [det.get("correct", "")]
        line += [r.score, r.max_score, round(r.percent, 2),
                 (r.result.entered_by.username if r.result and r.result.entered_by else "")]
        out.append(line)
    return headers, out


def block_ranking_table(subjects: Sequence[Subject], rows, aggregation: str) -> tuple[list, list]:
    """Blok reytingi jadvali: (sarlavhalar, qatorlar)."""
    headers = ["O'rin", "ID", "Familiya", "Ism", "Maktab", "Sinf"]
    headers += [s.name for s in subjects]
    headers += ["Jami ball", "Maks. ball", "Foiz" if aggregation == "sum" else "O'rtacha foiz"]
    out = []
    for r in rows:
        p = r.participant
        line = [r.place, p.code, p.last_name, p.first_name,
                p.school.name if p.school else "", p.grade.name if p.grade else ""]
        for s in subjects:
            if s.id not in r.parts:
                line.append("—")       # bu fanni tanlamagan
            elif r.parts[s.id] is None:
                line.append(0)         # tanlagan, lekin natija yo'q
            else:
                line.append(r.parts[s.id].score)
        line += [r.score, r.max_score, round(r.percent, 2)]
        out.append(line)
    return headers, out


def subject_ranking_xlsx(title: str, subject: Subject, rows, show_details: bool = True) -> bytes:
    wb = Workbook()
    headers, out = subject_ranking_table(subject, rows, show_details)
    _sheet(wb, title, headers, out, widths=[7, 8, 18, 16, 26, 8], first=True, places_col=1)
    return _to_bytes(wb)


def block_ranking_xlsx(title: str, subjects: Sequence[Subject], rows, aggregation: str) -> bytes:
    wb = Workbook()
    headers, out = block_ranking_table(subjects, rows, aggregation)
    _sheet(wb, title, headers, out, widths=[7, 8, 18, 16, 26, 8], first=True, places_col=1)
    return _to_bytes(wb)


def workbook_from_sheets(sheets: Sequence[dict]) -> bytes:
    """Bir nechta varaqli Excel. sheets: [{"title", "headers", "rows", "widths"?, "places"?: bool}]."""
    wb = Workbook()
    used: set[str] = set()
    for i, sh in enumerate(sheets):
        # Varaq nomlari takrorlanmasligi va 31 belgidan oshmasligi kerak
        base = "".join(c for c in sh["title"] if c not in '[]:*?/\\')[:28] or "Varaq"
        title, n = base, 2
        while title.lower() in used:
            title, n = f"{base[:26]}~{n}", n + 1
        used.add(title.lower())
        _sheet(wb, title, sh["headers"], sh["rows"], widths=sh.get("widths"), first=(i == 0),
               places_col=1 if sh.get("places") else None)
    return _to_bytes(wb)
